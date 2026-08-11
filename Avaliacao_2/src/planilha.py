"""Extrai cadastros dos PDFs aprovados e atualiza a planilha mestra."""

from __future__ import annotations

import re
import shutil
import unicodedata
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook
from pypdf import PdfReader


def _texto_pdf(caminho: Path) -> str:
    return "\n".join(
        pagina.extract_text() or "" for pagina in PdfReader(caminho).pages
    )


def _campo(texto: str, rotulo: str) -> str:
    resultado = re.search(
        rf"{re.escape(rotulo)}\s*:\s*(?:\r?\n\s*)?([^\r\n]+)",
        texto,
        re.IGNORECASE,
    )
    return resultado.group(1).strip() if resultado else ""


def _normalizar(texto: object) -> str:
    return "".join(
        caractere
        for caractere in unicodedata.normalize("NFKD", str(texto)).casefold()
        if not unicodedata.combining(caractere)
    ).strip()


def _cadastro_vazio(nome: str) -> dict[str, object]:
    return {
        "CPF": "",
        "Nome": nome,
        "Data de Nascimento": "",
        "Endereço": "",
        "E-mail": "",
        "Telefone": "",
        "Status": "OK",
        "Data de Processamento": datetime.now(),
        "Observações": "Dados extraídos automaticamente dos PDFs.",
    }


def extrair_dados(diretorio: Path) -> list[dict[str, object]]:
    """Agrupa documentos pelo nome e devolve um cadastro para cada pessoa."""
    cadastros: dict[str, dict[str, object]] = {}

    for arquivo in sorted(diretorio.glob("*.pdf")):
        texto = _texto_pdf(arquivo)
        nome = _campo(texto, "Nome") or _campo(texto, "Cliente")
        if not nome:
            continue
        chave = _normalizar(nome)
        dados = cadastros.setdefault(chave, _cadastro_vazio(nome))

        cpf = re.search(r"\b\d{3}\.?\d{3}\.?\d{3}-?\d{2}\b", texto)
        if cpf:
            dados["CPF"] = cpf.group()

        nascimento = _campo(texto, "Data de nascimento") or _campo(texto, "Nascimento")
        if nascimento:
            try:
                dados["Data de Nascimento"] = datetime.strptime(
                    nascimento, "%d/%m/%Y"
                ).date()
            except ValueError:
                dados["Data de Nascimento"] = nascimento

        endereco = _campo(texto, "Endereco") or _campo(texto, "Endereço")
        if endereco:
            complemento = [
                _campo(texto, "Bairro"),
                _campo(texto, "Cidade/UF"),
                _campo(texto, "CEP ficticio") or _campo(texto, "CEP"),
            ]
            dados["Endereço"] = ", ".join([endereco, *filter(None, complemento)])

    obrigatorios = ("CPF", "Nome", "Data de Nascimento", "Endereço")
    completos: list[dict[str, object]] = []
    for dados in cadastros.values():
        ausentes = [campo for campo in obrigatorios if not dados[campo]]
        if ausentes:
            print(
                f"Cadastro ignorado ({dados['Nome']}): faltam {', '.join(ausentes)}."
            )
        else:
            completos.append(dados)
    return completos


def atualizar_planilha(planilha: Path, diretorio_documentos: Path) -> list[int]:
    cadastros = extrair_dados(diretorio_documentos)
    if not cadastros:
        raise ValueError("Nenhum cadastro completo foi extraído dos PDFs.")

    workbook = load_workbook(planilha)
    aba = workbook["Planilha_Mestra"]
    cabecalhos = {celula.value: celula.column for celula in aba[1] if celula.value}
    linhas_atualizadas: list[int] = []

    for dados in cadastros:
        linha_destino = None
        primeira_vazia = None
        for linha in range(2, aba.max_row + 1):
            nome_existente = aba.cell(linha, cabecalhos["Nome"]).value
            cpf_existente = aba.cell(linha, cabecalhos["CPF"]).value
            if (
                nome_existente
                and _normalizar(nome_existente) == _normalizar(dados["Nome"])
                and re.sub(r"\D", "", str(cpf_existente or ""))
                == re.sub(r"\D", "", str(dados["CPF"]))
            ):
                linha_destino = linha
                break
            if primeira_vazia is None and not any(
                aba.cell(linha, coluna).value
                for coluna in range(1, aba.max_column + 1)
            ):
                primeira_vazia = linha

        linha_destino = linha_destino or primeira_vazia or aba.max_row + 1
        for cabecalho, valor in dados.items():
            if cabecalho in cabecalhos:
                aba.cell(linha_destino, cabecalhos[cabecalho], valor)
        aba.cell(linha_destino, cabecalhos["CPF"]).number_format = "@"
        aba.cell(linha_destino, cabecalhos["Data de Nascimento"]).number_format = (
            "dd/mm/yyyy"
        )
        aba.cell(linha_destino, cabecalhos["Data de Processamento"]).number_format = (
            "dd/mm/yyyy hh:mm:ss"
        )
        linhas_atualizadas.append(linha_destino)

    try:
        workbook.save(planilha)
    except PermissionError as erro:
        raise PermissionError(
            f"Não foi possível salvar {planilha.name}. Feche a planilha no Excel "
            "ou em outro visualizador e execute novamente."
        ) from erro

    diretorio_arquivados = diretorio_documentos.parent / "Documentos_Arquivados"
    diretorio_arquivados.mkdir(parents=True, exist_ok=True)
    for arquivo in sorted(diretorio_documentos.glob("*.pdf")):
        destino = diretorio_arquivados / arquivo.name
        contador = 1
        while destino.exists():
            destino = diretorio_arquivados / (
                f"{arquivo.stem}_{contador}{arquivo.suffix}"
            )
            contador += 1
        shutil.move(str(arquivo), str(destino))
        print(f"Documento arquivado: {destino}")

    return linhas_atualizadas


if __name__ == "__main__":
    raiz = Path(__file__).resolve().parents[1]
    try:
        linhas = atualizar_planilha(
            raiz / "Planilha_Mestra.xlsx", raiz / "Documentos_OK"
        )
        print(f"Planilha mestra atualizada nas linhas: {', '.join(map(str, linhas))}.")
    except (PermissionError, ValueError) as erro:
        raise SystemExit(str(erro)) from erro
